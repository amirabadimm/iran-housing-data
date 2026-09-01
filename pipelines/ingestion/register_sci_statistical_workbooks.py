"""Register the two manually downloaded SCI statistical workbooks.

The operation is byte-preserving. Run without ``--apply`` to inspect the move
plan; use ``--apply`` only after the two expected workbooks have been reviewed.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
INCOMING = ROOT / "data" / "incoming"
REGISTRY = ROOT / "metadata" / "file_registry.csv"
FIELDS = [
    "file_id", "dataset_id", "original_name", "relative_path",
    "publication_date", "downloaded_at", "checksum_sha256", "bytes",
    "inspection_status",
]
SOURCES = {
    "ts_building_140404-14050319105758.xlsx": {
        "dataset_id": "sci_tehran_residential_building_input_prices",
        "destination": "data/raw/construction_costs/input_price_indices/sci_tehran_residential_building_input_prices/1404/sci_tehran_residential_building_input_prices_1404_q4.xlsx",
        "expected_sheets": 22,
    },
    "ts_urban_140504-14050519161936.xlsx": {
        "dataset_id": "sci_urban_consumer_price_index",
        "destination": "data/raw/macroeconomic_environment/inflation/sci_urban_consumer_price_index/1405/sci_urban_consumer_price_index_1405_04.xlsx",
        "expected_sheets": 14,
    },
}


def digest(path: Path) -> str:
    result = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            result.update(chunk)
    return result.hexdigest()


def inspect(path: Path, expected_sheets: int) -> None:
    workbook = load_workbook(path, read_only=False, data_only=False)
    if len(workbook.sheetnames) != expected_sheets:
        raise ValueError(
            f"{path.name}: expected {expected_sheets} sheets, found "
            f"{len(workbook.sheetnames)}"
        )
    required = {"فهرست", "فراداده"}
    if not required.issubset(workbook.sheetnames):
        raise ValueError(f"{path.name}: missing required sheets {required}")


def plan() -> list[dict[str, object]]:
    result = []
    for original_name, config in SOURCES.items():
        source = INCOMING / original_name
        destination = ROOT / str(config["destination"])
        if not source.exists():
            if destination.exists():
                continue
            raise FileNotFoundError(source)
        if destination.exists():
            raise FileExistsError(destination)
        inspect(source, int(config["expected_sheets"]))
        result.append({
            **config,
            "source": source,
            "destination_path": destination,
            "checksum": digest(source),
            "bytes": source.stat().st_size,
            "downloaded_at": datetime.fromtimestamp(source.stat().st_mtime)
            .astimezone().isoformat(timespec="seconds"),
        })
    return result


def register(items: list[dict[str, object]]) -> None:
    with REGISTRY.open("r", encoding="utf-8-sig", newline="") as source:
        rows = list(csv.DictReader(source))
    known = {row["checksum_sha256"] for row in rows}
    for item in items:
        checksum = str(item["checksum"])
        if checksum in known:
            raise ValueError(f"Checksum already registered: {checksum}")
        source = Path(item["source"])
        destination = Path(item["destination_path"])
        rows.append({
            "file_id": f"sha256:{checksum}",
            "dataset_id": str(item["dataset_id"]),
            "original_name": source.name,
            "relative_path": destination.relative_to(ROOT).as_posix(),
            "publication_date": "",
            "downloaded_at": str(item["downloaded_at"]),
            "checksum_sha256": checksum,
            "bytes": str(item["bytes"]),
            "inspection_status": "inspected_xlsx_readable_complete_sheet_inventory",
        })
    temporary = REGISTRY.with_suffix(".csv.tmp")
    with temporary.open("w", encoding="utf-8", newline="") as target:
        writer = csv.DictWriter(target, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerows(rows)
    temporary.replace(REGISTRY)


def apply(items: list[dict[str, object]]) -> None:
    for item in items:
        source = Path(item["source"])
        destination = Path(item["destination_path"])
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(source, destination)
        if digest(destination) != item["checksum"]:
            raise OSError(f"Integrity check failed: {destination}")
    register(items)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()
    items = plan()
    if not items:
        print("Both SCI workbooks are already registered.")
        return
    for item in items:
        print(
            f"{Path(item['source']).relative_to(ROOT).as_posix()} -> "
            f"{Path(item['destination_path']).relative_to(ROOT).as_posix()}"
        )
    print(f"Validated {len(items)} SCI workbooks.")
    if args.apply:
        apply(items)
        print(f"Registered {len(items)} byte-preserved workbooks.")
    else:
        print("Dry run only; pass --apply to move and register them.")


if __name__ == "__main__":
    main()
