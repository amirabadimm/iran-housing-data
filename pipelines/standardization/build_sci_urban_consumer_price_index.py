"""Build a normalized long-form SCI urban consumer-price-index CSV."""

from __future__ import annotations

import argparse
import csv
import hashlib
import re
import unicodedata
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
DATASET_ID = "sci_urban_consumer_price_index"
RAW_ROOT = ROOT / "data" / "raw" / "macroeconomic_environment" / "inflation" / DATASET_ID
OUTPUT = ROOT / "data" / "standardized" / "macroeconomic_environment" / "inflation" / DATASET_ID / "urban_cpi_series.csv"
MONTHS = {name: index for index, name in enumerate(
    ["فروردین", "اردیبهشت", "خرداد", "تیر", "مرداد", "شهریور", "مهر", "آبان", "آذر", "دی", "بهمن", "اسفند"], 1
)}
METRICS = {
    "جدول 1": "price_index",
    "جدول 2": "month_over_month_percent_change",
    "جدول 3": "year_over_year_percent_change",
    "جدول 4": "twelve_month_inflation_percent",
    "جدول 5": "price_index",
    "جدول 6": "annual_percent_change",
    "جدول 7": "price_index",
    "جدول 8": "month_over_month_percent_change",
    "جدول 9": "year_over_year_percent_change",
    "جدول 10": "twelve_month_inflation_percent",
    "جدول 11": "price_index",
    "جدول 12": "price_index",
}


def text(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", str(value))).strip().translate(
        str.maketrans("يىك", "ییک")
    )


def number(value: object) -> float | None:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    value = text(value).replace(",", "")
    if value in {"", "-", "×", "*"}:
        return None
    try:
        return float(value)
    except ValueError:
        return None


def year_number(value: object) -> int | None:
    parsed = number(value)
    if parsed is not None and 1300 <= parsed <= 1499:
        return int(parsed)
    match = re.search(r"(?<!\d)(1[34]\d{2})(?!\d)", text(value))
    return int(match.group(1)) if match else None


def digest(path: Path) -> str:
    result = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            result.update(chunk)
    return result.hexdigest()


def source() -> Path:
    files = list(RAW_ROOT.rglob("*.xlsx"))
    if len(files) != 1:
        raise ValueError(f"Expected one raw workbook, found {len(files)}")
    return files[0]


def period_columns(sheet: object, start_column: int) -> tuple[list[tuple[int, int, int | None]], int]:
    year_row = max(range(1, 6), key=lambda row: sum(year_number(sheet.cell(row, column).value) is not None for column in range(start_column, sheet.max_column + 1)))
    period_row = max(range(1, 6), key=lambda row: sum(text(sheet.cell(row, column).value) in MONTHS or "\u0633\u0627\u0644\u0627\u0646\u0647" in text(sheet.cell(row, column).value) or text(sheet.cell(row, column).value) == "\u0633\u0627\u0644" for column in range(start_column, sheet.max_column + 1)))
    current_year = None
    result = []
    for column in range(start_column, sheet.max_column + 1):
        year_value = year_number(sheet.cell(year_row, column).value)
        if year_value is not None:
            current_year = year_value
        label = text(sheet.cell(period_row, column).value)
        if current_year is None:
            continue
        if label in MONTHS:
            result.append((column, current_year, MONTHS[label]))
        elif "سالانه" in label or label == "سال":
            result.append((column, current_year, None))
    return result, max(year_row, period_row) + 1


def add_row(
    rows: list[dict[str, object]], path: Path, checksum: str, table: str,
    metric: str, scope: str, geography_code: str, geography_fa: str,
    category_fa: str, year: int, month: int | None, value: float,
    source_column: int = 0,
    source_label_year: int | None = None,
    period_label_status: str = "as_published",
) -> None:
    rows.append({
        "metric": metric,
        "base_solar_hijri_year": 1400,
        "series_scope": scope,
        "geography_code": geography_code,
        "geography_fa": geography_fa,
        "category_fa": category_fa,
        "frequency": "monthly" if month else "annual",
        "solar_hijri_year": year,
        "solar_hijri_month": month or "",
        "solar_hijri_period": f"{year:04d}-{month:02d}" if month else f"{year:04d}",
        "source_label_year": source_label_year if source_label_year is not None else year,
        "period_label_status": period_label_status,
        "value": value,
        "unit": "index" if metric == "price_index" else "percent",
        "source_table": table.replace("جدول ", "table_"),
        "source_column": source_column,
        "source_file": path.relative_to(ROOT).as_posix(),
        "source_sha256": checksum,
    })


def extract(path: Path) -> list[dict[str, object]]:
    workbook = load_workbook(path, data_only=True, read_only=False)
    checksum = digest(path)
    rows: list[dict[str, object]] = []
    for table_number in range(1, 11):
        table = f"جدول {table_number}"
        sheet = workbook[table]
        provincial = table_number >= 7
        columns, data_start = period_columns(sheet, 3 if provincial else 2)
        if not columns:
            raise ValueError(f"No period columns in {table}")
        for row_index in range(data_start, sheet.max_row + 1):
            if provincial:
                geography_code = text(sheet.cell(row_index, 1).value)
                geography = text(sheet.cell(row_index, 2).value)
                category = "شاخص کل"
            else:
                geography_code, geography = "99", "کل کشور - خانوارهای شهری"
                category = text(sheet.cell(row_index, 1).value)
            if not geography or not category:
                continue
            for column, year, month in columns:
                value = number(sheet.cell(row_index, column).value)
                if value is not None:
                    source_label_year = year
                    period_label_status = "as_published"
                    if table_number == 9 and 159 <= column <= 170:
                        year = 1403
                        period_label_status = "corrected_source_label_1402_to_1403"
                    add_row(rows, path, checksum, table, METRICS[table],
                            "province" if provincial else "national_major_group",
                            geography_code, geography, category, year, month, value, column,
                            source_label_year, period_label_status)

    sheet = workbook["جدول 11"]
    for row_index in range(3, sheet.max_row + 1):
        year = number(sheet.cell(row_index, 1).value)
        value = number(sheet.cell(row_index, 2).value)
        if year is not None and value is not None:
            add_row(rows, path, checksum, "جدول 11", METRICS["جدول 11"],
                    "national_headline_historical", "99", "کل کشور - خانوارهای شهری",
                    "شاخص کل", int(year), None, value)

    sheet = workbook["جدول 12"]
    current_year = None
    for row_index in range(3, sheet.max_row + 1):
        year = number(sheet.cell(row_index, 1).value)
        if year is not None:
            current_year = int(year)
        month = MONTHS.get(text(sheet.cell(row_index, 2).value))
        value = number(sheet.cell(row_index, 3).value)
        if current_year is not None and month and value is not None:
            add_row(rows, path, checksum, "جدول 12", METRICS["جدول 12"],
                    "national_headline_historical", "99", "کل کشور - خانوارهای شهری",
                    "شاخص کل", current_year, month, value)
    return rows


def validate(rows: list[dict[str, object]]) -> None:
    if not rows:
        raise ValueError("No CPI records extracted")
    keys = [
        (r["source_table"], r["metric"], r["series_scope"], r["geography_code"],
         r["category_fa"], r["solar_hijri_period"], r["source_column"])
        for r in rows
    ]
    if len(keys) != len(set(keys)):
        raise ValueError("Duplicate CPI keys")
    if not {f"table_{n}" for n in range(1, 13)}.issubset({r["source_table"] for r in rows}):
        raise ValueError("At least one source table produced no observations")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--check", action="store_true")
    args = parser.parse_args()
    rows = extract(source())
    validate(rows)
    if not args.check:
        OUTPUT.parent.mkdir(parents=True, exist_ok=True)
        temporary = OUTPUT.with_suffix(".csv.tmp")
        with temporary.open("w", encoding="utf-8", newline="") as target:
            writer = csv.DictWriter(target, fieldnames=list(rows[0]))
            writer.writeheader()
            writer.writerows(rows)
        temporary.replace(OUTPUT)
    print(f"Validated {len(rows)} urban CPI observations from all 12 tables.")


if __name__ == "__main__":
    main()
