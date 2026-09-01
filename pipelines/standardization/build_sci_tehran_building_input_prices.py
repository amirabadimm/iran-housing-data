"""Build normalized SCI Tehran building-input price CSV files."""

from __future__ import annotations

import argparse
import csv
import hashlib
import re
import unicodedata
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
DATASET_ID = "sci_tehran_residential_building_input_prices"
RAW_ROOT = ROOT / "data" / "raw" / "construction_costs" / "input_price_indices" / DATASET_ID
OUTPUT_ROOT = ROOT / "data" / "standardized" / "construction_costs" / "input_price_indices" / DATASET_ID
QUARTERS = {"بهار": 1, "تابستان": 2, "پاییز": 3, "پاييز": 3, "زمستان": 4}
METRICS = {
    "جدول 1": ("price_index", 1402, "quarterly"),
    "جدول 2": ("quarter_over_quarter_percent_change", 1402, "quarterly"),
    "جدول 3": ("year_over_year_percent_change", 1402, "quarterly"),
    "جدول 4": ("four_quarter_inflation_percent", 1402, "quarterly"),
    "جدول 5": ("price_index", 1402, "annual"),
    "جدول 6": ("annual_percent_change", 1402, "annual"),
    "جدول7": ("price_index", 1390, "quarterly"),
}


def text(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", str(value))).strip().translate(
        str.maketrans("يىك", "ییک")
    )


def number(value: object) -> float | None:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    value = text(value).replace(",", "")
    if value in {"", "-", "×", "*"}:
        return None
    try:
        return float(value)
    except ValueError:
        return None


def digest(path: Path) -> str:
    result = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            result.update(chunk)
    return result.hexdigest()


def source() -> Path:
    files = list(RAW_ROOT.rglob("*.xlsx"))
    if len(files) != 1:
        raise ValueError(f"Expected one raw workbook, found {len(files)}")
    return files[0]


def series_rows(workbook: object, path: Path) -> list[dict[str, object]]:
    rows = []
    checksum = digest(path)
    for sheet_name, (metric, base_year, frequency) in METRICS.items():
        sheet = workbook[sheet_name]
        year_row = max(range(1, 7), key=lambda row: sum((candidate := number(sheet.cell(row, column).value)) is not None and 1300 <= candidate <= 1499 for column in range(2, sheet.max_column + 1)))
        labels = QUARTERS if frequency == "quarterly" else {"\u0633\u0627\u0644": 1}
        period_row = max(range(1, 7), key=lambda row: sum(text(sheet.cell(row, column).value) in labels for column in range(2, sheet.max_column + 1)))
        data_start = max(year_row, period_row) + 1
        current_year = None
        columns = []
        for column in range(2, sheet.max_column + 1):
            candidate = number(sheet.cell(year_row, column).value)
            if candidate is not None and 1300 <= candidate <= 1499:
                current_year = int(candidate)
            period_label = text(sheet.cell(period_row, column).value)
            if current_year is None:
                continue
            if frequency == "quarterly" and period_label in QUARTERS:
                columns.append((column, current_year, QUARTERS[period_label], period_label))
            elif frequency == "annual" and period_label == "سال":
                columns.append((column, current_year, None, period_label))
        if not columns:
            raise ValueError(f"No period columns found in {sheet_name}")
        for row_index in range(data_start, sheet.max_row + 1):
            category = text(sheet.cell(row_index, 2).value)
            if not category:
                continue
            for column, year, quarter, _ in columns:
                value = number(sheet.cell(row_index, column).value)
                if value is None:
                    continue
                rows.append({
                    "metric": metric,
                    "base_solar_hijri_year": base_year,
                    "frequency": frequency,
                    "solar_hijri_year": year,
                    "solar_hijri_quarter": quarter or "",
                    "solar_hijri_period": f"{year:04d}-Q{quarter}" if quarter else f"{year:04d}",
                    "category_fa": category,
                    "value": value,
                    "unit": "index" if metric == "price_index" else "percent",
                    "geography": "Tehran city",
                    "source_table": sheet_name.replace("جدول", "table_"),
                    "source_file": path.relative_to(ROOT).as_posix(),
                    "source_sha256": checksum,
                })
    return rows


def material_rows(workbook: object, path: Path) -> list[dict[str, object]]:
    rows = []
    checksum = digest(path)
    for table_number in range(8, 21):
        candidates = [f"جدول {table_number}", f"جدول{table_number}"]
        sheet_name = next(name for name in candidates if name in workbook.sheetnames)
        sheet = workbook[sheet_name]
        title = text(sheet.cell(1, 2).value or sheet.cell(1, 1).value)
        group = re.sub(r"^.*گروه\s*", "", title).split(":", 1)[0].strip()
        for row_index in range(4, sheet.max_row + 1):
            item = text(sheet.cell(row_index, 2).value)
            unit = text(sheet.cell(row_index, 3).value)
            minimum = number(sheet.cell(row_index, 4).value)
            maximum = number(sheet.cell(row_index, 5).value)
            average = number(sheet.cell(row_index, 6).value)
            if not item:
                continue
            rows.append({
                "solar_hijri_period": "1404-Q4",
                "solar_hijri_year": 1404,
                "solar_hijri_quarter": 4,
                "material_group_fa": group,
                "item_fa": item,
                "transaction_unit_fa": unit,
                "minimum_price_rials": minimum if minimum is not None else "",
                "maximum_price_rials": maximum if maximum is not None else "",
                "average_price_rials": average if average is not None else "",
                "geography": "Tehran city",
                "source_table": f"table_{table_number}",
                "source_row": row_index,
                "source_file": path.relative_to(ROOT).as_posix(),
                "source_sha256": checksum,
            })
    return rows


def validate(series: list[dict[str, object]], materials: list[dict[str, object]]) -> None:
    if not series or not materials:
        raise ValueError("Both series and material-price outputs must be non-empty")
    keys = [
        (row["source_table"], row["metric"], row["solar_hijri_period"], row["category_fa"])
        for row in series
    ]
    if len(keys) != len(set(keys)):
        raise ValueError("Duplicate building-input series keys")
    material_keys = [(row["source_table"], row["source_row"]) for row in materials]
    if len(material_keys) != len(set(material_keys)):
        raise ValueError("Duplicate material-price keys")


def write_csv(path: Path, rows: list[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(".csv.tmp")
    with temporary.open("w", encoding="utf-8", newline="") as target:
        writer = csv.DictWriter(target, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    temporary.replace(path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--check", action="store_true")
    args = parser.parse_args()
    path = source()
    workbook = load_workbook(path, data_only=True, read_only=False)
    series = series_rows(workbook, path)
    materials = material_rows(workbook, path)
    validate(series, materials)
    if not args.check:
        write_csv(OUTPUT_ROOT / "building_input_price_series.csv", series)
        write_csv(OUTPUT_ROOT / "building_material_prices.csv", materials)
    print(f"Validated {len(series)} index/inflation rows and {len(materials)} material-price rows.")


if __name__ == "__main__":
    main()
