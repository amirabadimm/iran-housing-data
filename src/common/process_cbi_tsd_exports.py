"""Register and standardize manually collected CBI TSD Excel exports.

Raw workbooks are copied byte-for-byte. Cleaned outputs retain source units and
values, restrict observations to Solar Hijri year 1370 onward, and expose CBI's
gray-cell preliminary status. Geographic series are written as one urban/Tehran
pair per indicator; a missing counterpart is never invented.
"""

from __future__ import annotations

import csv
import hashlib
import re
import shutil
from dataclasses import dataclass
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[2]
INCOMING = ROOT / "data" / "incoming" / "manually_collected"
SOURCE_URL = "https://tsdview.cis.cbi.ir/single-data"
COLLECTED_DATE = "2026-07-22"
REPORT_DATE = "1405/04/31"
MACRO_CLEAN_DIRECTORIES = {
    "liquidity": "money_and_credit",
    "employment": "labor_market",
    "national_accounts": "national_accounts",
}


@dataclass(frozen=True)
class ColumnRef:
    filename: str
    column: int


@dataclass(frozen=True)
class PairSpec:
    dataset_id: str
    name: str
    unit: str
    urban: ColumnRef
    tehran: ColumnRef
    subcategory: str


@dataclass(frozen=True)
class SingleSpec:
    dataset_id: str
    name: str
    unit: str
    refs: tuple[tuple[str, ColumnRef], ...]
    category: str
    subcategory: str
    collected_date: str = COLLECTED_DATE
    related_task: str = "cbi_tsd_bulk_standardization_20260722"


PAIR_SPECS = (
    PairSpec("cbi_hsg_private_investment_new_buildings_q", "Private-sector investment in new urban buildings", "billion rial", ColumnRef("cbi_private_building_investment_urban_quarterly.xlsx", 2), ColumnRef("cbi_private_building_investment_tehran_quarterly.xlsx", 2), "construction_investment"),
    PairSpec("cbi_hsg_started_buildings_count_q", "Private-sector started buildings: count", "building", ColumnRef("cbi_private_buildings_started_count_urban_quarterly.xlsx", 2), ColumnRef("cbi_private_buildings_started_count_tehran_quarterly.xlsx", 2), "construction_starts"),
    PairSpec("cbi_hsg_started_buildings_floor_area_q", "Private-sector started buildings: total floor area", "thousand square metres", ColumnRef("cbi_private_buildings_started_floor_area_urban_quarterly.xlsx", 2), ColumnRef("cbi_private_buildings_started_floor_area_tehran_quarterly.xlsx", 2), "construction_starts"),
    PairSpec("cbi_hsg_started_buildings_cost_per_sqm_q", "Private-sector started buildings: estimated cost per square metre", "thousand rial", ColumnRef("cbi_private_buildings_started_cost_per_sqm_urban_tehran_quarterly.xlsx", 3), ColumnRef("cbi_private_buildings_started_cost_per_sqm_urban_tehran_quarterly.xlsx", 2), "construction_cost"),
    PairSpec("cbi_hsg_started_buildings_land_value_per_sqm_q", "Private-sector started buildings: land value per square metre", "thousand rial", ColumnRef("cbi_private_buildings_land_value_per_sqm_started_completed_urban_tehran_quarterly.xlsx", 2), ColumnRef("cbi_private_buildings_land_value_per_sqm_started_completed_urban_tehran_quarterly.xlsx", 5), "land_value"),
    PairSpec("cbi_hsg_completed_buildings_land_value_per_sqm_q", "Private-sector completed buildings: land value per square metre", "thousand rial", ColumnRef("cbi_private_buildings_land_value_per_sqm_started_completed_urban_tehran_quarterly.xlsx", 4), ColumnRef("cbi_private_buildings_land_value_per_sqm_started_completed_urban_tehran_quarterly.xlsx", 3), "land_value"),
    PairSpec("cbi_hsg_completed_buildings_floor_area_q", "Private-sector completed buildings: total floor area", "thousand square metres", ColumnRef("cbi_private_buildings_completed_floor_area_cost_urban_tehran_quarterly.xlsx", 3), ColumnRef("cbi_private_buildings_completed_floor_area_cost_urban_tehran_quarterly.xlsx", 2), "construction_completions"),
    PairSpec("cbi_hsg_completed_buildings_cost_per_sqm_q", "Private-sector completed buildings: estimated cost per square metre", "thousand rial", ColumnRef("cbi_private_buildings_completed_floor_area_cost_urban_tehran_quarterly.xlsx", 5), ColumnRef("cbi_private_buildings_completed_floor_area_cost_urban_tehran_quarterly.xlsx", 4), "construction_cost"),
    PairSpec("cbi_hsg_completed_buildings_count_q", "Private-sector completed buildings: count", "building", ColumnRef("cbi_private_buildings_completed_count_urban_tehran_quarterly.xlsx", 3), ColumnRef("cbi_private_buildings_completed_count_urban_tehran_quarterly.xlsx", 2), "construction_completions"),
    PairSpec("cbi_hsg_building_permits_count_q", "Building permits issued by urban municipalities: count", "permit", ColumnRef("cbi_residential_units_completed_and_building_permits_count_quarterly.xlsx", 4), ColumnRef("cbi_residential_units_completed_and_building_permits_count_quarterly.xlsx", 3), "construction_permits"),
    PairSpec("cbi_hsg_rent_index_q", "Rental housing index", "index", ColumnRef("cbi_rent_land_price_unemployment_indices_quarterly.xlsx", 7), ColumnRef("cbi_rent_land_price_unemployment_indices_quarterly.xlsx", 2), "rent"),
    PairSpec("cbi_hsg_land_price_index_q", "Land price index", "index", ColumnRef("cbi_rent_land_price_unemployment_indices_quarterly.xlsx", 3), ColumnRef("cbi_rent_land_price_unemployment_indices_quarterly.xlsx", 8), "land_price"),
    PairSpec("cbi_hsg_building_permits_floor_area_q", "Building permits issued by urban municipalities: estimated floor area", "thousand square metres", ColumnRef("cbi_building_permits_floor_area_quarterly.xlsx", 2), ColumnRef("cbi_building_permits_floor_area_quarterly.xlsx", 4), "construction_permits"),
)

SINGLE_SPECS = (
    SingleSpec("cbi_hsg_private_sector_credit_q", "Outstanding bank and credit-institution facilities to the private housing/construction sector", "thousand billion rial", (("value", ColumnRef("cbi_private_sector_housing_credit_quarterly.xlsx", 2)),), "housing", "housing_credit"),
    SingleSpec("cbi_hsg_private_investment_total_q", "Private-sector investment in new urban buildings: total by construction stage", "billion rial", (("value", ColumnRef("cbi_private_building_investment_by_construction_stage_quarterly.xlsx", 2)),), "housing", "construction_investment"),
    SingleSpec("cbi_hsg_completed_residential_units_floor_area_q", "Private-sector completed residential units: total floor area", "thousand square metres", (("value", ColumnRef("cbi_completed_residential_units_floor_area_quarterly.xlsx", 2)),), "housing", "construction_completions"),
    SingleSpec("cbi_hsg_completed_residential_units_average_area_q", "Private-sector completed residential units: average floor area", "square metres", (("value", ColumnRef("cbi_completed_residential_units_floor_area_quarterly.xlsx", 3)),), "housing", "construction_completions"),
    SingleSpec("cbi_hsg_completed_residential_units_count_q", "Private-sector completed residential units: count", "unit", (("value", ColumnRef("cbi_residential_units_completed_and_building_permits_count_quarterly.xlsx", 2)),), "housing", "construction_completions"),
    SingleSpec("cbi_hsg_rent_index_by_city_size_q", "Rental housing index by city-size group", "index", (("large_cities_value", ColumnRef("cbi_rent_land_price_unemployment_indices_quarterly.xlsx", 4)), ("medium_cities_value", ColumnRef("cbi_rent_land_price_unemployment_indices_quarterly.xlsx", 5)), ("small_cities_value", ColumnRef("cbi_rent_land_price_unemployment_indices_quarterly.xlsx", 6))), "housing", "rent"),
    SingleSpec("cbi_hsg_bank_maskan_loans_a", "Bank Maskan loans paid", "mixed; see column units", (("count_thousand_loans", ColumnRef("cbi_bank_maskan_loans_annual.xlsx", 2)), ("amount_billion_rial", ColumnRef("cbi_bank_maskan_loans_annual.xlsx", 3))), "housing", "mortgage_credit"),
    SingleSpec("cbi_macro_liquidity_q", "Liquidity", "thousand billion rial", (("value", ColumnRef("cbi_liquidity_by_components_quarterly_1379_1404.xlsx", 2)),), "macro", "liquidity"),
    SingleSpec("cbi_macro_urban_unemployment_rate_q", "Urban unemployment rate", "percent", (("value", ColumnRef("cbi_rent_land_price_unemployment_indices_quarterly.xlsx", 9)),), "macro", "employment"),
    SingleSpec("cbi_macro_building_gfcf_private_current_a", "Private-sector gross fixed capital formation in buildings at current prices", "billion rial", (("value", ColumnRef("cbi_building_investment_real_estate_value_added_national_accounts_annual_1395_1402.xlsx", 2)),), "macro", "national_accounts", "2026-07-23", "cbi_tsd_national_accounts_20260723"),
    SingleSpec("cbi_macro_building_gfcf_public_current_a", "Public-sector gross fixed capital formation in buildings at current prices", "billion rial", (("value", ColumnRef("cbi_building_investment_real_estate_value_added_national_accounts_annual_1395_1402.xlsx", 3)),), "macro", "national_accounts", "2026-07-23", "cbi_tsd_national_accounts_20260723"),
    SingleSpec("cbi_macro_real_estate_value_added_current_a", "Real-estate activities value added at current prices", "billion rial", (("value", ColumnRef("cbi_building_investment_real_estate_value_added_national_accounts_annual_1395_1402.xlsx", 4)),), "macro", "national_accounts", "2026-07-23", "cbi_tsd_national_accounts_20260723"),
    SingleSpec("cbi_macro_building_gfcf_private_constant_1400_a", "Private-sector gross fixed capital formation in buildings at constant 1400 prices", "billion rial", (("value", ColumnRef("cbi_building_investment_real_estate_value_added_national_accounts_annual_1395_1402.xlsx", 7)),), "macro", "national_accounts", "2026-07-23", "cbi_tsd_national_accounts_20260723"),
    SingleSpec("cbi_macro_building_gfcf_public_constant_1400_a", "Public-sector gross fixed capital formation in buildings at constant 1400 prices", "billion rial", (("value", ColumnRef("cbi_building_investment_real_estate_value_added_national_accounts_annual_1395_1402.xlsx", 8)),), "macro", "national_accounts", "2026-07-23", "cbi_tsd_national_accounts_20260723"),
    SingleSpec("cbi_macro_real_estate_value_added_constant_1400_a", "Real-estate activities value added at constant 1400 prices", "billion rial", (("value", ColumnRef("cbi_building_investment_real_estate_value_added_national_accounts_annual_1395_1402.xlsx", 12)),), "macro", "national_accounts", "2026-07-23", "cbi_tsd_national_accounts_20260723"),
)


def checksum(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def period_parts(value: object) -> tuple[str, int, int | None]:
    text = str(value).strip()
    year_match = re.match(r"^(\d{4})", text)
    if not year_match:
        raise ValueError(f"Unrecognized period: {text!r}")
    year = int(year_match.group(1))
    quarter = None
    quarter_words = {"اول": 1, "دوم": 2, "سوم": 3, "چهارم": 4}
    for word, number in quarter_words.items():
        if word in text:
            quarter = number
            break
    normalized = f"{year}-Q{quarter}" if quarter else str(year)
    return normalized, year, quarter


def is_preliminary(cell: openpyxl.cell.cell.Cell) -> bool:
    return cell.fill.fill_type == "solid" and cell.fill.fgColor.rgb in {"00D9D9D9", "FFD9D9D9"}


def read_ref(ref: ColumnRef) -> tuple[dict[str, tuple[object, bool, int, int | None]], dict[str, str]]:
    path = INCOMING / ref.filename
    workbook = openpyxl.load_workbook(path, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    observations: dict[str, tuple[object, bool, int, int | None]] = {}
    for row in range(9, sheet.max_row + 1):
        raw_period = sheet.cell(row, 1).value
        value = sheet.cell(row, ref.column).value
        if raw_period is None or value is None:
            continue
        period, year, quarter = period_parts(raw_period)
        if year < 1370:
            continue
        observations[period] = (value, is_preliminary(sheet.cell(row, ref.column)), year, quarter)
    metadata = {
        "frequency_fa": str(sheet["B4"].value),
        "source_description_fa": str(sheet.cell(7, ref.column).value),
        "source_label_fa": str(sheet.cell(6, ref.column).value),
        "source_unit_fa": str(sheet.cell(8, ref.column).value),
    }
    return observations, metadata


def period_key(period: str) -> tuple[int, int]:
    match = re.match(r"^(\d{4})(?:-Q(\d))?$", period)
    assert match
    return int(match.group(1)), int(match.group(2) or 0)


def write_csv(path: Path, fieldnames: list[str], rows: list[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def read_csv_rows(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def relative(path: Path) -> str:
    return path.relative_to(ROOT).as_posix()


def copy_raw_files() -> dict[str, Path]:
    result: dict[str, Path] = {}
    for source in sorted(INCOMING.glob("*.xlsx")):
        if source.name == "cbi_building_investment_real_estate_value_added_national_accounts_annual_1395_1402.xlsx":
            category, batch = "macro", "cbi_tsd_14050501"
        else:
            category = "macro" if source.name == "cbi_liquidity_by_components_quarterly_1379_1404.xlsx" else "housing"
            batch = "cbi_tsd_14050431"
        destination = ROOT / "data" / "raw" / category / batch / source.name
        destination.parent.mkdir(parents=True, exist_ok=True)
        if destination.exists() and checksum(destination) != checksum(source):
            raise FileExistsError(f"Raw destination differs and will not be overwritten: {destination}")
        if not destination.exists():
            shutil.copy2(source, destination)
        result[source.name] = destination
    return result


def catalog_row(dataset_id: str, name: str, category: str, subcategory: str, unit: str, output: Path, refs: list[ColumnRef], observations: list[dict[str, object]], raw_paths: dict[str, Path], notes: str, collected_date: str = COLLECTED_DATE, related_task: str = "cbi_tsd_bulk_standardization_20260722") -> dict[str, str]:
    periods = [str(row["period"]) for row in observations]
    source_files = sorted({ref.filename for ref in refs})
    return {
        "dataset_id": dataset_id,
        "dataset_name": name,
        "topic": "Central Bank of Iran time-series data",
        "category": category,
        "subcategory": subcategory,
        "category_status": "confirmed",
        "description": name,
        "source_organization": "Central Bank of the Islamic Republic of Iran",
        "source_url": SOURCE_URL,
        "collection_method": "manual Excel export",
        "date_collected": collected_date,
        "original_filename": "; ".join(source_files),
        "stored_filename": "; ".join(path.name for path in (raw_paths[name] for name in source_files)),
        "raw_path": "; ".join(relative(raw_paths[name]) for name in source_files),
        "cleaned_path": relative(output),
        "derived_path": "",
        "file_format": "CSV",
        "sheet_names": "فصلي" if dataset_id.endswith("_q") else "سالانه",
        "time_coverage_start": min(periods, key=period_key) if periods else "",
        "time_coverage_end": max(periods, key=period_key) if periods else "",
        "frequency": "quarterly" if dataset_id.endswith("_q") else "annual",
        "geographic_coverage": "all urban areas and Tehran" if len(refs) == 2 and dataset_id not in {"cbi_hsg_bank_maskan_loans_a"} else "as labeled by CBI; see notes",
        "unit": unit,
        "language": "Persian source; English standardized fields",
        "access_status": "public website",
        "cleaning_status": "completed",
        "validation_status": "automated structural checks completed; substantive review recommended",
        "current_use": "available",
        "related_task": related_task,
        "confidentiality": "public",
        "checksum_sha256": checksum(output),
        "notes": notes,
    }


def main() -> None:
    raw_paths = copy_raw_files()
    catalog: list[dict[str, str]] = []
    cleaning_log: list[dict[str, object]] = []
    variable_rows: list[dict[str, str]] = []

    for spec in PAIR_SPECS:
        urban, urban_meta = read_ref(spec.urban)
        tehran, tehran_meta = read_ref(spec.tehran)
        periods = sorted(set(urban) | set(tehran), key=period_key)
        rows: list[dict[str, object]] = []
        for period in periods:
            sample = urban.get(period) or tehran.get(period)
            assert sample
            rows.append({
                "period": period,
                "solar_hijri_year": sample[2],
                "quarter": sample[3],
                "urban_value": urban.get(period, ("", False, 0, None))[0],
                "tehran_value": tehran.get(period, ("", False, 0, None))[0],
                "unit": spec.unit,
                "urban_preliminary": urban.get(period, (None, False, 0, None))[1],
                "tehran_preliminary": tehran.get(period, (None, False, 0, None))[1],
            })
        output = ROOT / "data" / "cleaned" / "housing" / "cbi_urban_tehran_pairs" / f"{spec.dataset_id}.csv"
        fields = ["period", "solar_hijri_year", "quarter", "urban_value", "tehran_value", "unit", "urban_preliminary", "tehran_preliminary"]
        refs = [spec.urban, spec.tehran]
        source_name = urban_meta["source_description_fa"]
        source_unit = urban_meta["source_unit_fa"]
        for row in rows:
            row["unit"] = source_unit
        write_csv(output, fields, rows)
        note = f"نام و واحد مستقیماً از فایل اکسل حفظ شده‌اند. برچسب‌ها: urban={urban_meta['source_label_fa']}; Tehran={tehran_meta['source_label_fa']}. Urban and Tehran values are paired by period; no interpolation; blanks represent unavailable counterparts."
        catalog.append(catalog_row(spec.dataset_id, source_name, "housing", spec.subcategory, source_unit, output, refs, rows, raw_paths, note))
        for variable in ("urban_value", "tehran_value"):
            meta = urban_meta if variable == "urban_value" else tehran_meta
            variable_rows.append({"dataset_id": spec.dataset_id, "variable_name": variable, "variable_label": meta["source_label_fa"], "description": meta["source_description_fa"], "data_type": "numeric", "unit": meta["source_unit_fa"], "language": "Persian", "allowed_values": "", "missing_value_codes": "blank", "source_definition": meta["source_description_fa"], "notes": "Name, label, unit, and values are preserved directly from the CBI Excel export."})
        cleaning_log.append({"cleaning_id": f"clean_{spec.dataset_id}", "dataset_id": spec.dataset_id, "date": COLLECTED_DATE, "input_path": "; ".join(relative(raw_paths[r.filename]) for r in refs), "output_path": relative(output), "script": "src/common/process_cbi_tsd_exports.py", "transformation": "Extracted source columns; paired urban and Tehran by normalized Solar Hijri quarter; retained years >=1370; preserved values, units, missingness, and preliminary flags.", "reason": "User-requested standardization", "rows_before": len(set(read_ref(spec.urban)[0]) | set(read_ref(spec.tehran)[0])), "rows_after": len(rows), "columns_before": 2, "columns_after": len(fields), "performed_by": "Codex", "review_status": "completed", "notes": "No imputation, aggregation, or unit conversion."})

    for spec in SINGLE_SPECS:
        extracted = [(label, *read_ref(ref), ref) for label, ref in spec.refs]
        periods = sorted(set().union(*(set(values) for _, values, _, _ in extracted)), key=period_key)
        rows = []
        for period in periods:
            sample = next(values[period] for _, values, _, _ in extracted if period in values)
            row: dict[str, object] = {"period": period, "solar_hijri_year": sample[2]}
            if sample[3] is not None:
                row["quarter"] = sample[3]
            for label, values, _, _ in extracted:
                row[label] = values.get(period, ("", False, 0, None))[0]
                row[f"{label}_preliminary"] = values.get(period, (None, False, 0, None))[1]
            rows.append(row)
        clean_directory = (
            MACRO_CLEAN_DIRECTORIES[spec.subcategory]
            if spec.category == "macro"
            else "cbi_non_geographic"
        )
        output = ROOT / "data" / "cleaned" / spec.category / clean_directory / f"{spec.dataset_id}.csv"
        fields = ["period", "solar_hijri_year"] + (["quarter"] if spec.dataset_id.endswith("_q") else [])
        for label, _, _, _ in extracted:
            fields.extend([label, f"{label}_preliminary"])
        write_csv(output, fields, rows)
        refs = [ref for _, _, _, ref in extracted]
        source_names = list(dict.fromkeys(meta["source_description_fa"] for _, _, meta, _ in extracted))
        source_units = list(dict.fromkeys(meta["source_unit_fa"] for _, _, meta, _ in extracted))
        source_name = " | ".join(source_names)
        source_unit = " | ".join(source_units)
        note = "نام‌ها، برچسب‌ها و واحدها مستقیماً از فایل اکسل حفظ شده‌اند. This source series is not an urban–Tehran geographic pair and no counterpart was invented. Source labels: " + "; ".join(meta["source_label_fa"] for _, _, meta, _ in extracted)
        catalog.append(catalog_row(spec.dataset_id, source_name, spec.category, spec.subcategory, source_unit, output, refs, rows, raw_paths, note, spec.collected_date, spec.related_task))
        for label, _, meta, _ in extracted:
            variable_rows.append({"dataset_id": spec.dataset_id, "variable_name": label, "variable_label": meta["source_label_fa"], "description": meta["source_description_fa"], "data_type": "numeric", "unit": meta["source_unit_fa"], "language": "Persian", "allowed_values": "", "missing_value_codes": "blank", "source_definition": meta["source_description_fa"], "notes": "Name, label, unit, and values are preserved directly from the CBI Excel export."})
        cleaning_log.append({"cleaning_id": f"clean_{spec.dataset_id}", "dataset_id": spec.dataset_id, "date": spec.collected_date, "input_path": "; ".join(relative(raw_paths[r.filename]) for r in refs), "output_path": relative(output), "script": "src/common/process_cbi_tsd_exports.py", "transformation": "Extracted source series; normalized Solar Hijri period; retained years >=1370; preserved values, missingness, source units, and preliminary flags.", "reason": "User-requested standardization", "rows_before": len(periods), "rows_after": len(rows), "columns_before": len(refs), "columns_after": len(fields), "performed_by": "Codex", "review_status": "completed", "notes": "No geographic counterpart existed in the supplied files; none was fabricated."})

    catalog = [row for row in read_csv_rows(ROOT / "metadata" / "data_catalog.csv") if not row.get("dataset_id", "").startswith("cbi_")] + catalog
    variable_rows = [row for row in read_csv_rows(ROOT / "metadata" / "variable_dictionary.csv") if not row.get("dataset_id", "").startswith("cbi_")] + variable_rows
    cleaning_log = [row for row in read_csv_rows(ROOT / "metadata" / "cleaning_log.csv") if not row.get("dataset_id", "").startswith("cbi_")] + cleaning_log

    catalog_fields = "dataset_id,dataset_name,topic,category,subcategory,category_status,description,source_organization,source_url,collection_method,date_collected,original_filename,stored_filename,raw_path,cleaned_path,derived_path,file_format,sheet_names,time_coverage_start,time_coverage_end,frequency,geographic_coverage,unit,language,access_status,cleaning_status,validation_status,current_use,related_task,confidentiality,checksum_sha256,notes".split(",")
    write_csv(ROOT / "metadata" / "data_catalog.csv", catalog_fields, catalog)
    variable_fields = "dataset_id,variable_name,variable_label,description,data_type,unit,language,allowed_values,missing_value_codes,source_definition,notes".split(",")
    write_csv(ROOT / "metadata" / "variable_dictionary.csv", variable_fields, variable_rows)
    cleaning_fields = "cleaning_id,dataset_id,date,input_path,output_path,script,transformation,reason,rows_before,rows_after,columns_before,columns_after,performed_by,review_status,notes".split(",")
    write_csv(ROOT / "metadata" / "cleaning_log.csv", cleaning_fields, cleaning_log)
    source_rows = [row for row in read_csv_rows(ROOT / "metadata" / "source_registry.csv") if row.get("source_id") != "cbi_tsd"]
    source_rows.append({"source_id": "cbi_tsd", "source_organization": "Central Bank of the Islamic Republic of Iran", "source_name": "Time Series Database (TSD)", "source_url": SOURCE_URL, "access_method": "manual Excel export", "date_accessed": "2026-07-23", "license": "not stated in supplied workbooks; requires review", "access_status": "public website", "contact": "", "notes": f"19 Excel exports: 18 generated {REPORT_DATE} and one annual national-accounts export generated 1405/05/01; website timed out during automated verification on {COLLECTED_DATE}."})
    write_csv(ROOT / "metadata" / "source_registry.csv", "source_id,source_organization,source_name,source_url,access_method,date_accessed,license,access_status,contact,notes".split(","), source_rows)

    inventory_rows: list[dict[str, object]] = []
    for source in sorted(INCOMING.glob("*.xlsx")):
        workbook = openpyxl.load_workbook(source, data_only=True)
        for sheet in workbook.worksheets:
            for column in range(2, sheet.max_column + 1):
                populated = []
                preliminary_count = 0
                for row in range(9, sheet.max_row + 1):
                    value = sheet.cell(row, column).value
                    if value is not None:
                        populated.append((sheet.cell(row, 1).value, value))
                        preliminary_count += int(is_preliminary(sheet.cell(row, column)))
                inventory_rows.append({
                    "source_file": source.name,
                    "sheet_name": sheet.title,
                    "report_title_fa": sheet["A1"].value,
                    "report_date_solar_hijri": sheet["B2"].value,
                    "reported_range_fa": sheet["B3"].value,
                    "frequency_fa": sheet["B4"].value,
                    "column_number": column,
                    "source_label_fa": sheet.cell(6, column).value,
                    "source_dataset_name_fa": sheet.cell(7, column).value,
                    "source_unit_fa": sheet.cell(8, column).value,
                    "observation_count": len(populated),
                    "first_source_period": populated[0][0] if populated else "",
                    "last_source_period": populated[-1][0] if populated else "",
                    "preliminary_observation_count": preliminary_count,
                    "status": "available" if populated else "empty_source_column",
                })
    inventory_fields = ["source_file", "sheet_name", "report_title_fa", "report_date_solar_hijri", "reported_range_fa", "frequency_fa", "column_number", "source_label_fa", "source_dataset_name_fa", "source_unit_fa", "observation_count", "first_source_period", "last_source_period", "preliminary_observation_count", "status"]
    write_csv(ROOT / "metadata" / "excel_series_inventory.csv", inventory_fields, inventory_rows)

    issues = [row for row in read_csv_rows(ROOT / "metadata" / "data_issues.csv") if not row.get("issue_id", "").startswith("issue_cbi_")] + [
        {"issue_id": "issue_cbi_empty_construction_services_index", "dataset_id": "", "date_identified": COLLECTED_DATE, "issue_type": "empty_source_series", "description": "Column J in cbi_rent_land_price_unemployment_indices_quarterly.xlsx is labeled construction services price index (1400=100) but contains no observations.", "severity": "medium", "status": "open", "resolution": "Re-export this series from CBI TSD if it is required.", "related_file": "data/raw/housing/cbi_tsd_14050431/cbi_rent_land_price_unemployment_indices_quarterly.xlsx", "notes": "Not registered as an available cleaned dataset because there are zero values."},
        {"issue_id": "issue_cbi_empty_bank_maskan_quarterly_count", "dataset_id": "", "date_identified": COLLECTED_DATE, "issue_type": "empty_source_series", "description": "Column C in cbi_building_permits_floor_area_quarterly.xlsx is labeled as a Bank Maskan loan-count series but has no unit and no observations.", "severity": "medium", "status": "open", "resolution": "Re-export the quarterly series from CBI TSD if it is required; the separate annual export remains available.", "related_file": "data/raw/housing/cbi_tsd_14050431/cbi_building_permits_floor_area_quarterly.xlsx", "notes": "Not registered as an available cleaned dataset because there are zero values."},
        {"issue_id": "issue_cbi_empty_national_accounts_columns", "dataset_id": "", "date_identified": "2026-07-23", "issue_type": "empty_source_series", "description": "Six labeled columns (E, F, I, J, K, and M) in cbi_building_investment_real_estate_value_added_national_accounts_annual_1395_1402.xlsx contain no observations.", "severity": "low", "status": "open", "resolution": "Re-export the specific CBI TSD series if they are required and the portal supplies observations.", "related_file": "data/raw/macro/cbi_tsd_14050501/cbi_building_investment_real_estate_value_added_national_accounts_annual_1395_1402.xlsx", "notes": "All columns remain documented in excel_series_inventory.csv; only six populated columns were registered as cleaned datasets."},
    ]
    write_csv(ROOT / "metadata" / "data_issues.csv", "issue_id,dataset_id,date_identified,issue_type,description,severity,status,resolution,related_file,notes".split(","), issues)

    manifest_rows: list[dict[str, object]] = []
    for layer in ("raw", "cleaned"):
        for path in sorted((ROOT / "data" / layer).glob("**/*")):
            if not path.is_file() or path.name in {".gitkeep", "README.md"}:
                continue
            manifest_rows.append({
                "layer": layer,
                "category": path.relative_to(ROOT / "data" / layer).parts[0],
                "path": relative(path),
                "file_format": path.suffix.lower().lstrip("."),
                "size_bytes": path.stat().st_size,
                "checksum_sha256": checksum(path),
            })
    write_csv(
        ROOT / "metadata" / "file_manifest.csv",
        ["layer", "category", "path", "file_format", "size_bytes", "checksum_sha256"],
        manifest_rows,
    )

    assert sum(row["dataset_id"].startswith("cbi_") for row in catalog) == 28
    assert len(list((ROOT / "data" / "raw").glob("**/*.xlsx"))) >= 19
    assert all(row["time_coverage_start"] >= "1370" for row in catalog if row["dataset_id"].startswith("cbi_"))
    assert len(manifest_rows) >= 40
    print("Registered and cleaned 28 CBI datasets from 19 raw workbooks; preserved other registered datasets.")


if __name__ == "__main__":
    main()
